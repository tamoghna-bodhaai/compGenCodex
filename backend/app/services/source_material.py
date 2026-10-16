"""Deterministic answer-key mapping and source lifecycle."""
from __future__ import annotations

import io
import json
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path

from docx import Document

from app.db.database import get_connection
from app.schemas.assessment import ResponseSpec, legacy_response_spec, validate_answer, rendered_stem, generator_compatible
from app.services.diagrams import asset_path
from app.schemas.ingestion import ClassifiedSeedQuestion
from app.services.seed_import import upsert_questions


def _now() -> str:
    return datetime.now(UTC).isoformat()


def attachment_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt":
            return content.decode("utf-8-sig")
        if suffix == ".docx":
            document = Document(io.BytesIO(content))
            return "\n".join([p.text for p in document.paragraphs] + [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        if suffix == ".pdf":
            from app.services.ingestion import _best_pdf_pages
            return "\n".join(text for _, text in _best_pdf_pages(content))
    except Exception as error:
        raise ValueError("This answer-key or solution file could not be read.") from error
    raise ValueError("Answer-key uploads must be TXT, PDF, or DOCX.")


def parse_entries(text: str) -> list[dict]:
    entries = []
    blocks: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if re.match(r"^(?:Q(?:uestion)?\s*)?\d+\s*[.):\-]", line, re.I) or not blocks:
            blocks.append(line)
        else:
            blocks[-1] += "\n" + line
    for line in blocks:
        match = re.match(r"^(?:Q(?:uestion)?\s*)?(\d+)\s*[.):\-]\s*(.+)$", line, re.I | re.S)
        if not match:
            for value in line.splitlines():
                if re.fullmatch(r"(?:[A-D]|[-+]?\d+(?:\.\d+)?)", value, re.I):
                    entries.append({"number": None, "answer": value, "solution": None, "raw": value})
            continue
        body = match.group(2).strip()
        pair_text = re.match(r"^((?:[A-Za-z][A-Za-z0-9_]*\s*(?:->|:|-)\s*[A-Za-z0-9_]+\s*[,;]\s*)+[A-Za-z][A-Za-z0-9_]*\s*(?:->|:|-)\s*[A-Za-z0-9_]+)(.*)$", body, re.S)
        if pair_text:
            pairs = re.findall(r"([A-Za-z][A-Za-z0-9_]*)\s*(?:->|:|-)\s*([A-Za-z0-9_]+)", pair_text.group(1))
            entries.append({"number": int(match.group(1)), "answer": dict(pairs), "solution": pair_text.group(2).strip(" ;:-") or None, "raw": line})
            continue
        answer_match = re.match(r"^(?:Ans(?:wer)?\s*[:=]\s*)?((?:[A-D](?:\s*[,/&+]\s*[A-D])*)|[-+]?\d+(?:\.\d+)?)\b(.*)$", body, re.I | re.S)
        answer = answer_match.group(1).strip() if answer_match else None
        solution = answer_match.group(2).strip(" ;:-") if answer_match and answer_match.group(2).strip(" ;:-") else (body if not answer else None)
        entries.append({"number": int(match.group(1)), "answer": answer, "solution": solution, "raw": line})
    return entries


def add_attachment(job_id: str, kind: str, *, filename: str = "pasted.txt", content: bytes = b"", text: str = "", map_now: bool = True) -> dict:
    if kind not in {"answer_key", "solutions"}:
        raise ValueError("Unsupported attachment kind.")
    source = text.strip() or attachment_text(filename, content)
    identifier, now = str(uuid.uuid4()), _now()
    with get_connection() as connection:
        if not connection.execute("SELECT 1 FROM ingestion_jobs WHERE id = ?", (job_id,)).fetchone():
            raise LookupError("Ingestion job not found.")
        connection.execute("INSERT INTO source_attachments (id, job_id, kind, filename, content, text_content, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                           (identifier, job_id, kind, filename, content or None, source, now))
    return map_attachment(job_id, identifier) if map_now else {"attachment_id": identifier, "mappings": []}


def map_attachment(job_id: str, attachment_id: str) -> dict:
    with get_connection() as connection:
        attachment = connection.execute("SELECT * FROM source_attachments WHERE id = ? AND job_id = ?", (attachment_id, job_id)).fetchone()
        if attachment is None:
            raise LookupError("Source attachment not found.")
        entries = parse_entries(attachment["text_content"] or "")
        questions = connection.execute("SELECT * FROM questions WHERE source_job_id = ?", (job_id,)).fetchall()
        ordered = sorted(questions, key=lambda row: (row["source_question_number"] is None, row["source_question_number"] or 0))
        order_is_unambiguous = bool(entries) and all(entry["number"] is None for entry in entries) and len(entries) == len(ordered) and len({row["source_question_number"] for row in ordered}) == len(ordered)
        numbers: dict[int, list] = {}
        for row in questions:
            if row["source_question_number"] is not None:
                numbers.setdefault(row["source_question_number"], []).append(row)
        results = []
        for position, entry in enumerate(entries):
            matches = [ordered[position]] if order_is_unambiguous else numbers.get(entry["number"], [])
            answer = entry["answer"]
            if isinstance(answer, str) and re.fullmatch(r"[A-D](?:\s*[,/&+]\s*[A-D])+", answer, re.I):
                answer = re.split(r"\s*[,/&+]\s*", answer.upper())
            target = matches[0] if len(matches) == 1 else None
            status = "matched" if target else ("ambiguous" if matches else "unmatched")
            if target and answer is not None:
                spec = ResponseSpec.model_validate(json.loads(target["response_spec_json"])) if target["response_spec_json"] else legacy_response_spec(target["question_type"], len(json.loads(target["question_json"]).get("options", [])))
                if not validate_answer(spec, answer, len(json.loads(target["question_json"]).get("options", []))):
                    status = "ambiguous"
            mapping_id = str(uuid.uuid4())
            connection.execute("INSERT INTO source_mappings (id, job_id, attachment_id, source_question_number, question_id, answer_json, solution, status, note, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                               (mapping_id, job_id, attachment_id, entry["number"], target["id"] if target else None, json.dumps({"correct_answer": answer}) if answer else None, entry["solution"], status, entry["raw"], _now()))
            if status == "matched" and target:
                if answer is not None:
                    connection.execute("UPDATE questions SET answer_json = ?, answer_provenance = 'source', updated_at = ? WHERE id = ?", (json.dumps({"correct_answer": answer}), _now(), target["id"]))
                if entry["solution"]:
                    connection.execute("UPDATE questions SET solution = ?, solution_provenance = 'source', updated_at = ? WHERE id = ?", (entry["solution"], _now(), target["id"]))
            results.append({"id": mapping_id, "number": entry["number"], "question_id": target["id"] if target else None, "status": status})
    return {"attachment_id": attachment_id, "mappings": results}


def delete_question(question_id: str) -> None:
    with get_connection() as connection:
        rows = connection.execute("SELECT * FROM question_diagrams WHERE seed_question_id = ?", (question_id,)).fetchall()
        if not connection.execute("SELECT 1 FROM questions WHERE id = ?", (question_id,)).fetchone():
            raise LookupError("Question not found.")
        job_row = connection.execute("SELECT source_job_id FROM questions WHERE id = ?", (question_id,)).fetchone()
        connection.execute("UPDATE ingestion_candidates SET promoted_question_id = NULL, status = 'deleted' WHERE promoted_question_id = ?", (question_id,))
        connection.execute("DELETE FROM questions WHERE id = ?", (question_id,))
        if job_row and job_row["source_job_id"]:
            connection.execute("UPDATE ingestion_jobs SET accepted_questions = (SELECT COUNT(*) FROM ingestion_candidates WHERE job_id = ? AND status = 'accepted') WHERE id = ?", (job_row["source_job_id"], job_row["source_job_id"]))
    for row in rows:
        path = asset_path(dict(row))
        if path:
            path.unlink(missing_ok=True)


def delete_source(job_id: str) -> None:
    with get_connection() as connection:
        job = connection.execute("SELECT * FROM ingestion_jobs WHERE id = ?", (job_id,)).fetchone()
        if job is None:
            raise LookupError("Ingestion job not found.")
        if job["state"] in {"queued", "running"}:
            raise ValueError("An active source cannot be deleted.")
        ids = [row[0] for row in connection.execute("SELECT id FROM questions WHERE source_job_id = ?", (job_id,))]
    for identifier in ids:
        delete_question(identifier)
    with get_connection() as connection:
        connection.execute("INSERT INTO source_audit (id, source_name, question_count, deleted_at) VALUES (?, ?, ?, ?)", (job_id, job["source_name"], len(ids), _now()))
        connection.execute("DELETE FROM lifecycle_events WHERE job_id = ?", (job_id,))
        connection.execute("DELETE FROM llm_cost_ledger WHERE job_id = ?", (job_id,))
        connection.execute("DELETE FROM bank_solution_jobs WHERE source_job_id = ?", (job_id,))
        connection.execute("DELETE FROM ingestion_jobs WHERE id = ?", (job_id,))


def map_manually(mapping_id: str, question_id: str) -> dict:
    with get_connection() as connection:
        mapping = connection.execute("SELECT * FROM source_mappings WHERE id = ?", (mapping_id,)).fetchone()
        question = connection.execute("SELECT * FROM questions WHERE id = ?", (question_id,)).fetchone()
        if mapping is None or question is None or question["source_job_id"] != mapping["job_id"]:
            raise LookupError("Mapping or question not found in this source.")
        answer = json.loads(mapping["answer_json"])["correct_answer"] if mapping["answer_json"] else None
        spec = ResponseSpec.model_validate(json.loads(question["response_spec_json"])) if question["response_spec_json"] else legacy_response_spec(question["question_type"], len(json.loads(question["question_json"]).get("options", [])))
        if not validate_answer(spec, answer, len(json.loads(question["question_json"]).get("options", []))):
            raise ValueError("Answer does not fit this question's response rule.")
        if answer is not None:
            connection.execute("UPDATE questions SET answer_json = ?, answer_provenance = 'source', updated_at = ? WHERE id = ?", (mapping["answer_json"], _now(), question_id))
        if mapping["solution"]:
            connection.execute("UPDATE questions SET solution = ?, solution_provenance = 'source', updated_at = ? WHERE id = ?", (mapping["solution"], _now(), question_id))
        connection.execute("UPDATE source_mappings SET question_id = ?, status = 'matched' WHERE id = ?", (question_id, mapping_id))
    return {"id": mapping_id, "question_id": question_id, "status": "matched"}


def promote_candidate(job_id: str, candidate_id: str, changes: dict) -> dict:
    with get_connection() as connection:
        row = connection.execute("SELECT c.*, j.source_name FROM ingestion_candidates c JOIN ingestion_jobs j ON j.id = c.job_id WHERE c.id = ? AND c.job_id = ?", (candidate_id, job_id)).fetchone()
    if row is None:
        raise LookupError("Review candidate not found.")
    if row["status"] not in {"review", "accepted"}:
        raise ValueError("Candidate cannot be promoted.")
    payload = json.loads(row["payload_json"])
    payload.update(changes)
    classified = ClassifiedSeedQuestion.model_validate(payload)
    if len(classified.stem.strip()) < 12:
        raise ValueError("Enter a readable question stem before promotion.")
    spec = classified.response_spec or legacy_response_spec(classified.question_type.value, len(classified.options))
    spec.supported_for_generation = generator_compatible(spec, classified.question_type.value, len(classified.options))
    answer = classified.correct_answer
    if answer is not None and not validate_answer(spec, answer, len(classified.options)):
        raise ValueError("Answer does not fit the response rule.")
    if spec.mode == "unknown" or classified.format_uncertain:
        raise ValueError("Resolve the response format before promotion.")
    from app.services.ingestion import format_review_reason
    reason = format_review_reason(classified.model_dump(), spec)
    if reason:
        raise ValueError(reason)
    record = classified.model_dump()
    record.update({
        "source_key": f"ingested-{job_id[:12]}-review-{candidate_id[:12]}",
        "source_reference": f"{row['source_name']} - page {classified.source_page or 'not stated'}, question {classified.source_question_number}",
        "question_json": {"stem": rendered_stem(classified.stem, spec, classified.question_group), "options": classified.options},
        "answer_json": {"correct_answer": answer} if answer is not None else None,
        "source": row["source_name"], "source_job_id": job_id,
        "verification_status": "pending_review",
        "response_spec": spec.model_dump(),
        "answer_provenance": "source" if answer is not None else None,
        "solution_provenance": "source" if classified.solution else None,
    })
    upsert_questions([record])
    with get_connection() as connection:
        saved = connection.execute("SELECT id FROM questions WHERE source_key = ?", (record["source_key"],)).fetchone()
        connection.execute("UPDATE ingestion_candidates SET status = 'accepted', promoted_question_id = ?, payload_json = ?, validation_notes = 'Promoted after manual review', updated_at = ? WHERE id = ?",
                           (saved["id"], json.dumps(payload), _now(), candidate_id))
        connection.execute("UPDATE ingestion_jobs SET accepted_questions = (SELECT COUNT(*) FROM ingestion_candidates WHERE job_id = ? AND status = 'accepted'), review_questions = (SELECT COUNT(*) FROM ingestion_candidates WHERE job_id = ? AND status = 'review') WHERE id = ?", (job_id, job_id, job_id))
    return {"question_id": saved["id"], "status": "accepted"}
