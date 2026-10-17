from __future__ import annotations

import os
import io
import asyncio
import zipfile
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from unittest.mock import AsyncMock

from app.db.database import get_connection, decode_question_row, initialize_database
from app.schemas.assessment import MarkingRule, NumericRule, QuestionGroup, ResponseSpec, generator_compatible, legacy_response_spec, rendered_stem, validate_answer
from app.schemas.ingestion import ClassifiedSeedQuestion
from app.services.ingestion import QuestionIngestionService, format_review_reason
from app.services.ingestion import _render_docx_images_for_vision
from app.services.diagrams import crop_normalized_page, store_diagram, diagrams_for, asset_path
from app.services.seed_import import upsert_questions
from app.services.source_material import add_attachment, delete_source, delete_question, map_manually, parse_entries
from app.api.questions import question_catalog, export_ingested_source
from app.schemas.papers import ExportFormat, ExportVariant
from app.services.bank_solutions import queue_solution_job, run_solution_job, get_solution_job, latest_source_solution_job


class AssessmentCompatibilityTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.old_url = os.environ.get("DATABASE_URL")
        os.environ["DATABASE_URL"] = f"sqlite:///{Path(self.temporary.name) / 'test.db'}"

    def tearDown(self) -> None:
        if self.old_url is None:
            os.environ.pop("DATABASE_URL", None)
        else:
            os.environ["DATABASE_URL"] = self.old_url
        self.temporary.cleanup()

    def _seed(self, job_id: str, number: int, spec: ResponseSpec | None = None) -> str:
        key = f"source-q{number}"
        question = {
            "source_key": key, "exam": "JEE", "subject": "Mathematics", "chapter": "Algebra", "topic": "Matrices",
            "primary_concept": "Matrices", "question_archetype": "compute", "question_type": "single_correct_mcq",
            "difficulty": 3, "question_json": {"stem": "Choose the correct value of the expression.", "options": ["1", "2", "3", "4"]},
            "source": "paper.pdf", "source_reference": f"paper.pdf question {number}", "verification_status": "pending_review",
            "source_job_id": job_id, "source_question_number": number,
            "response_spec": spec.model_dump() if spec else None,
        }
        upsert_questions([question])
        with get_connection() as connection:
            return connection.execute("SELECT id FROM questions WHERE source_key = ?", (key,)).fetchone()[0]

    def test_response_rules_cover_main_and_advanced_variants(self) -> None:
        main_mcq = legacy_response_spec("single_correct_mcq", 4)
        self.assertTrue(validate_answer(main_mcq, "B", 4))
        self.assertFalse(validate_answer(main_mcq, ["A", "B"], 4))
        integer = ResponseSpec(mode="numeric", numeric=NumericRule(integer_only=True, permitted_precision=0, rounding_rule="nearest"))
        self.assertTrue(validate_answer(integer, "42", 0))
        self.assertFalse(validate_answer(integer, "42.5", 0))
        tolerant = ResponseSpec(mode="numeric", numeric=NumericRule(tolerance=0.01))
        self.assertTrue(validate_answer(tolerant, "2.005", expected="2.0"))
        self.assertFalse(validate_answer(tolerant, "2.02", expected="2.0"))
        multi = ResponseSpec(mode="select", min_selections=2, max_selections=3)
        self.assertTrue(validate_answer(multi, ["A", "C"], 4))
        self.assertFalse(validate_answer(multi, "A", 4))
        matching = ResponseSpec(mode="matching", list_i=["P", "Q"], list_ii=["1", "2"])
        self.assertTrue(validate_answer(matching, {"P": "2", "Q": "1"}))
        self.assertFalse(validate_answer(matching, {"P": "2"}))
        self.assertEqual(parse_entries("7. P-2, Q-1")[0]["answer"], {"P": "2", "Q": "1"})
        self.assertIn("Second step", parse_entries("1. Answer: B\nFirst step\nSecond step")[0]["solution"])
        comprehension = QuestionGroup(id="passage-1", type="comprehension", shared_stem="Read the common passage.")
        self.assertIn("Read the common passage.", rendered_stem("Find the value.", main_mcq, comprehension))
        self.assertIn("List I", rendered_stem("Match the following.", matching))
        paper_one = ResponseSpec(mode="select", min_selections=1, max_selections=1, instructions="Choose one.", marking=MarkingRule(full=4, negative=-1, source_instructions="+4, -1"))
        paper_two = ResponseSpec(mode="select", min_selections=1, max_selections=3, instructions="Choose all.", marking=MarkingRule(full=4, partial={"two correct": 2}, source_instructions="Partial credit applies"))
        self.assertNotEqual(paper_one.marking.model_dump(), paper_two.marking.model_dump())
        self.assertIsNone(format_review_reason({"exam_format": {"exam": "JEE Advanced"}, "response_spec": paper_two.model_dump()}, paper_two))
        self.assertIsNotNone(format_review_reason({"exam_format": {"exam": "JEE Advanced"}}, main_mcq))
        advanced_five = ClassifiedSeedQuestion.model_validate({"source_question_number": 1, "exam": "JEE", "subject": "Physics", "chapter": "Mechanics", "topic": "Motion", "primary_concept": "Motion", "question_archetype": "solve", "question_type": "single_correct_mcq", "difficulty": 3, "stem": "Select the correct physical result.", "options": ["A", "B", "C", "D", "E"], "response_spec": paper_one.model_dump()})
        self.assertEqual(len(advanced_five.options), 5)
        self.assertFalse(generator_compatible(paper_one.model_copy(update={"supported_for_generation": True}), "single_correct_mcq", 5))

    def test_legacy_question_gets_non_destructive_default(self) -> None:
        job = QuestionIngestionService().create_job("paper.pdf")
        identifier = self._seed(job["id"], 1)
        with get_connection() as connection:
            row = decode_question_row(connection.execute("SELECT * FROM questions WHERE id = ?", (identifier,)).fetchone())
        self.assertEqual(row["response_spec"]["mode"], "select")
        self.assertIsNone(row["answer_json"])

    def test_existing_ingestion_lineage_is_backfilled(self) -> None:
        job = QuestionIngestionService().create_job("old.pdf")
        identifier = self._seed(job["id"], 1)
        with get_connection() as connection:
            connection.execute("UPDATE questions SET source_job_id = NULL WHERE id = ?", (identifier,))
            connection.execute("UPDATE questions SET source_key = ? WHERE id = ?", (f"ingested-{job['id'][:12]}-q0001", identifier))
            connection.execute("DELETE FROM app_migrations WHERE name = 'backfill_ingestion_lineage_v1'")
        initialize_database()
        with get_connection() as connection:
            row = connection.execute("SELECT source_job_id, source_question_number FROM questions WHERE id = ?", (identifier,)).fetchone()
        self.assertEqual(tuple(row), (job["id"], 1))

    def test_generation_catalog_excludes_unsupported_format(self) -> None:
        job = QuestionIngestionService().create_job("advanced.pdf")
        self._seed(job["id"], 1, ResponseSpec(mode="matching", list_i=["P"], list_ii=["1"]))
        self.assertEqual(sum(row["count"] for row in question_catalog()["items"]), 1)
        self.assertEqual(question_catalog(supported_only=True)["items"], [])

    def test_source_export_contains_an_accepted_question(self) -> None:
        job = QuestionIngestionService().create_job("practice.pdf")
        self._seed(job["id"], 1)
        response = export_ingested_source(job["id"], ExportFormat.DOCX, ExportVariant.QUESTION_PAPER)
        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertGreater(len(response.body), 1000)

    def test_direct_mapping_and_source_deletion_keep_paper_snapshot(self) -> None:
        job = QuestionIngestionService().create_job("paper.pdf")
        first = self._seed(job["id"], 1)
        self._seed(job["id"], 2)
        with patch("app.services.openrouter.OpenRouterClient.call_llm", side_effect=AssertionError("LLM must not be called")):
            mapped = add_attachment(job["id"], "answer_key", text="1. B\n3. A")
        self.assertEqual([item["status"] for item in mapped["mappings"]], ["matched", "unmatched"])
        with get_connection() as connection:
            answer = connection.execute("SELECT answer_json FROM questions WHERE id = ?", (first,)).fetchone()[0]
            self.assertIn("B", answer)
        second_id = next(row[0] for row in get_rows("SELECT id FROM questions WHERE source_question_number = 2") )
        resolved = map_manually(mapped["mappings"][1]["id"], second_id)
        self.assertEqual(resolved["status"], "matched")
        with get_connection() as connection:
            connection.execute("UPDATE ingestion_jobs SET state = 'succeeded' WHERE id = ?", (job["id"],))
            connection.execute("INSERT INTO papers (id, title, exam, subject, generation_config, status, created_at, updated_at) VALUES ('paper', 'Saved', 'JEE', 'Mathematics', '{}', 'generated', 'now', 'now')")
            connection.execute("INSERT INTO paper_questions (id, paper_id, position, question_json, question_type, difficulty, created_at, updated_at) VALUES ('snapshot', 'paper', 1, '{\"stem\":\"Saved question\"}', 'single_correct_mcq', 3, 'now', 'now')")
        delete_source(job["id"])
        with get_connection() as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM questions").fetchone()[0], 0)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM source_attachments").fetchone()[0], 0)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM source_audit").fetchone()[0], 1)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM paper_questions").fetchone()[0], 1)

    def test_order_mapping_only_when_unambiguous(self) -> None:
        job = QuestionIngestionService().create_job("ordered.txt")
        self._seed(job["id"], 4)
        self._seed(job["id"], 5)
        mapped = add_attachment(job["id"], "answer_key", text="A\nC")
        self.assertEqual([item["status"] for item in mapped["mappings"]], ["matched", "matched"])
        ambiguous = add_attachment(job["id"], "answer_key", text="B")
        self.assertEqual(ambiguous["mappings"][0]["status"], "unmatched")

    def test_docx_embedded_image_can_be_associated_with_an_option(self) -> None:
        from docx import Document
        from PIL import Image
        document = Document()
        document.add_paragraph("Choose the diagram in option A.")
        picture = io.BytesIO()
        Image.new("RGB", (100, 80), "red").save(picture, format="PNG")
        picture.seek(0)
        document.add_picture(picture)
        output = io.BytesIO()
        document.save(output)
        sheet = _render_docx_images_for_vision(output.getvalue())
        self.assertIsNotNone(sheet)
        job = QuestionIngestionService().create_job("visual.docx")
        identifier = self._seed(job["id"], 1)
        crop = crop_normalized_page(sheet, [0, 0, 0.5, 0.5])
        diagram = store_diagram(owner_column="seed_question_id", owner_id=identifier, provenance="source_crop", data=crop, description="Option A figure", option_index=0)
        self.assertEqual(diagram["option_index"], 0)
        self.assertEqual(len(diagrams_for(owner_column="seed_question_id", owner_id=identifier)), 1)
        exported = export_ingested_source(job["id"], ExportFormat.DOCX, ExportVariant.QUESTION_PAPER)
        with zipfile.ZipFile(io.BytesIO(exported.body)) as archive:
            self.assertTrue(any(name.startswith("word/media/") for name in archive.namelist()))
            self.assertIn("Option A", archive.read("word/document.xml").decode())
        diagram_path = asset_path(diagram)
        delete_question(identifier)
        self.assertFalse(diagram_path.exists())

    def test_ai_solution_is_validated_and_marked_for_review(self) -> None:
        job = QuestionIngestionService().create_job("solutions.pdf")
        identifier = self._seed(job["id"], 1, legacy_response_spec("single_correct_mcq", 4))
        queued = queue_solution_job(source_job_id=job["id"], question_ids=[identifier])
        self.assertEqual(queued["items"], [{"question_id": identifier, "state": "queued", "error_message": None}])
        self.assertEqual(latest_source_solution_job(job["id"])["id"], queued["id"])
        self.assertTrue(queue_solution_job(source_job_id=job["id"], question_ids=[identifier])["reused"])
        fake = type("Result", (), {"correct_answer": "B", "solution": "Worked solution"})()
        async def generate(*args, **kwargs):
            self.assertEqual(get_solution_job(queued["id"])["items"][0]["state"], "running")
            return fake
        with patch("app.services.bank_solutions.GenerationService") as service:
            service.return_value.generate_solution = AsyncMock(side_effect=generate)
            asyncio.run(run_solution_job(queued["id"]))
        self.assertEqual(get_solution_job(queued["id"])["state"], "succeeded")
        self.assertEqual(get_solution_job(queued["id"])["items"][0]["state"], "succeeded")
        with get_connection() as connection:
            row = connection.execute("SELECT verification_status, answer_provenance, solution_provenance FROM questions WHERE id = ?", (identifier,)).fetchone()
        self.assertEqual(tuple(row), ("pending_review", "ai", "ai"))
        with get_connection() as connection:
            connection.execute("UPDATE questions SET answer_json = '{\"correct_answer\":\"B\"}', answer_provenance = 'source', solution = NULL WHERE id = ?", (identifier,))
        preserved = queue_solution_job(question_ids=[identifier])
        with patch("app.services.bank_solutions.GenerationService") as service:
            service.return_value.generate_solution = AsyncMock(return_value=fake)
            asyncio.run(run_solution_job(preserved["id"]))
        with get_connection() as connection:
            row = connection.execute("SELECT answer_provenance, solution_provenance FROM questions WHERE id = ?", (identifier,)).fetchone()
        self.assertEqual(tuple(row), ("source", "ai"))
        retry = queue_solution_job(question_ids=[identifier], replace_existing=True)
        fake.correct_answer = "E"
        with patch("app.services.bank_solutions.GenerationService") as service:
            service.return_value.generate_solution = AsyncMock(return_value=fake)
            asyncio.run(run_solution_job(retry["id"]))
        self.assertEqual(get_solution_job(retry["id"])["state"], "failed")
        self.assertEqual(get_solution_job(retry["id"])["items"][0]["state"], "failed")
        self.assertTrue(get_solution_job(retry["id"])["items"][0]["error_message"])

    def test_finished_solutions_are_visible_before_the_batch_finishes(self) -> None:
        source = QuestionIngestionService().create_job("batch.pdf")
        ids = [self._seed(source["id"], number, legacy_response_spec("single_correct_mcq", 4)) for number in (1, 2)]
        queued = queue_solution_job(source_job_id=source["id"])
        self.assertEqual([item["question_id"] for item in queued["items"]], ids)

        async def generate(*args, **kwargs):
            if get_solution_job(queued["id"])["items"][0]["state"] == "succeeded":
                progress = get_solution_job(queued["id"])
                self.assertEqual(progress["completed_questions"], 1)
                self.assertEqual([item["state"] for item in progress["items"]], ["succeeded", "running"])
                with get_connection() as connection:
                    self.assertEqual(connection.execute("SELECT solution FROM questions WHERE id = ?", (ids[0],)).fetchone()[0], "Step by step")
            return type("Result", (), {"correct_answer": "B", "solution": "Step by step"})()

        with patch("app.services.bank_solutions.GenerationService") as service:
            service.return_value.generate_solution = AsyncMock(side_effect=generate)
            asyncio.run(run_solution_job(queued["id"]))
        self.assertEqual(get_solution_job(queued["id"])["completed_questions"], 2)


def get_rows(query: str):
    with get_connection() as connection:
        return connection.execute(query).fetchall()
